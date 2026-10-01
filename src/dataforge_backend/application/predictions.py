"""Guided scoring: documented model versions, cohort scores with factor breakdowns, and exports."""

import io
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from ..core.errors import Conflict, NotFound, ValidationFailed
from ..core.security import utcnow
from ..domain.analysis import AnalysisDataset
from ..domain.analysis.stats import median, quantile, standard_deviation
from ..domain.prediction import (
    FactorInput,
    FactorSpec,
    ModelError,
    ModelSpec,
    factor_from_dict,
    score_cohort,
    validate_model,
)
from ..domain.roles import Permission
from ..infrastructure.db.models import ModelVersion, Prediction
from . import analysis, audit, catalogue
from .common import Actor, humanize


def model_spec(model: ModelVersion) -> ModelSpec:
    return ModelSpec(
        name=model.name,
        version=model.version,
        outcome=model.outcome,
        dataset_id=model.dataset_id,
        factors=tuple(factor_from_dict(f) for f in model.factors),
        description=model.description,
    )


def model_dict(model: ModelVersion) -> dict[str, Any]:
    return {
        "id": model.id,
        "name": model.name,
        "version": model.version,
        "label": f"{model.name} v{model.version}",
        "outcome": model.outcome,
        "description": model.description,
        "datasetId": model.dataset_id,
        "factors": model.factors,
        "weightsSum": round(sum(float(f["weight"]) for f in model.factors), 6),
        "status": model.status,
        "createdAt": model.created_at.isoformat(),
        "reviewedAt": model.reviewed_at.isoformat() if model.reviewed_at else None,
    }


async def list_models(session: AsyncSession, actor: Actor) -> list[dict[str, Any]]:
    actor.require(Permission.PREDICT)
    datasets = {d.id for d in await catalogue.analysis_catalogue(session, actor)}
    stmt = select(ModelVersion).order_by(ModelVersion.name, ModelVersion.created_at.desc())
    if not actor.can(Permission.CURATE_MODELS):
        stmt = stmt.where(ModelVersion.status == "active")
    return [model_dict(m) for m in (await session.execute(stmt)).scalars() if m.dataset_id in datasets]


async def create_version(session: AsyncSession, actor: Actor, body: dict[str, Any]) -> dict[str, Any]:
    """A new model version. Weights must add up to one and every factor must read a measure of the dataset."""
    actor.require(Permission.CURATE_MODELS, "Only a Model Curator can publish model versions.")
    dataset = await analysis.dataset_for(session, actor, str(body.get("datasetId", "")))
    try:
        factors = tuple(factor_from_dict(f) for f in body.get("factors") or [])
        spec = validate_model(
            ModelSpec(
                str(body.get("name", "")).strip(),
                str(body.get("version", "")).strip(),
                str(body.get("outcome", "")).strip(),
                dataset.id,
                factors,
                str(body.get("description", "")),
            ),
            set(dataset.measures),
        )
    except ModelError as error:
        raise ValidationFailed(str(error)) from error
    if not spec.name or not spec.version or not spec.outcome:
        raise ValidationFailed("Give the model a name, a version and the outcome it scores.")
    if (
        await session.execute(
            select(ModelVersion).where(ModelVersion.name == spec.name, ModelVersion.version == spec.version)
        )
    ).scalar_one_or_none():
        raise Conflict(f"{spec.label} already exists. Versions are immutable; publish a new version number.")
    if body.get("retirePrevious", True):
        for previous in (
            await session.execute(
                select(ModelVersion).where(ModelVersion.name == spec.name, ModelVersion.status == "active")
            )
        ).scalars():
            previous.status = "retired"
    model = ModelVersion(
        name=spec.name,
        version=spec.version,
        outcome=spec.outcome,
        description=spec.description,
        dataset_id=dataset.id,
        factors=[f.to_dict() for f in spec.factors],
        status="active",
        created_by=actor.id,
        reviewed_by=actor.id,
        reviewed_at=utcnow(),
    )
    session.add(model)
    await session.flush()
    await audit.record(
        session,
        actor,
        "Published model version",
        "model",
        model.id,
        spec.label,
        f"{len(factors)} factors, weights sum to 1.00",
    )
    await session.commit()
    return model_dict(model)


async def _factor_input(
    session: AsyncSession, actor: Actor, dataset: AnalysisDataset, factor: FactorSpec, filters: dict[str, str]
) -> FactorInput:
    months = list(dataset.months)
    measure = dataset.measures[factor.measure]
    dims = list(dataset.dimensions)
    recent = (months[-3], months[-1]) if len(months) >= 3 else (months[0], months[-1])
    before = (months[-6], months[-4]) if len(months) >= 6 else (months[0], months[max(0, len(months) - 4)])

    async def population_values(window: tuple[str, str]) -> dict[tuple[str, ...], float]:
        rows = await analysis.query_groups(
            session,
            actor,
            dataset,
            {"v": measure.id},
            dims,
            {},
            date_from=window[0],
            date_to=window[1],
            kind="prediction",
        )
        return {
            tuple(str(r[f"d{i}"]) for i in range(len(dims))): float(r["v"])
            for r in rows
            if r["v"] is not None
        }

    async def cohort_value(window: tuple[str, str]) -> float:
        rows = await analysis.query_groups(
            session,
            actor,
            dataset,
            {"v": measure.id},
            [],
            filters,
            date_from=window[0],
            date_to=window[1],
            kind="prediction",
        )
        return float(rows[0]["v"]) if rows and rows[0]["v"] is not None else float("nan")

    if factor.window == "trend":
        now, then = await population_values(recent), await population_values(before)
        # Each group's change, matched by the group's own key.
        population = [now[key] - then[key] for key in now if key in then]
        cohort = await cohort_value(recent) - await cohort_value(before)
    else:
        window = recent if factor.window == "recent" else (months[0], months[-1])
        population = list((await population_values(window)).values())
        cohort = await cohort_value(window)
    centre = median(population) if population else float("nan")
    iqr = quantile(population, 0.75) - quantile(population, 0.25) if population else 0.0
    # The interquartile range scaled to a standard deviation is robust to extreme groups.
    spread = iqr / 1.349 if iqr > 0 else (standard_deviation(population) if len(population) > 1 else 0.0)
    return FactorInput(cohort=cohort, centre=centre, spread=spread, units=len(population), unit=measure.unit)


def prediction_dict(prediction: Prediction, model: ModelVersion) -> dict[str, Any]:
    return {
        "id": prediction.id,
        "subject": prediction.subject,
        "outcome": model.outcome,
        "score": prediction.score,
        "band": prediction.band,
        "modelVersion": f"{model.name} v{model.version}",
        "modelId": model.id,
        "cohort": prediction.cohort,
        "factors": prediction.factors,
        "createdAt": prediction.created_at.isoformat(),
        "when": humanize(prediction.created_at),
    }


async def score(
    session: AsyncSession, actor: Actor, model_id: str | None, cohort: dict[str, str]
) -> dict[str, Any]:
    actor.require(Permission.PREDICT)
    actor.require_verified()
    if model_id:
        model = await session.get(ModelVersion, model_id)
    else:
        model = (
            (
                await session.execute(
                    select(ModelVersion)
                    .where(ModelVersion.status == "active")
                    .order_by(ModelVersion.created_at.desc())
                )
            )
            .scalars()
            .first()
        )
    if model is None:
        raise NotFound("No model is available. A Model Curator publishes one.")
    dataset = await analysis.dataset_for(session, actor, model.dataset_id)
    spec = model_spec(model)
    for dim, value in cohort.items():
        if dim not in dataset.dimensions or value not in dataset.dimensions[dim].values:
            raise ValidationFailed(f"{dim} = {value} is not a cohort of {dataset.name}.")
    if not dataset.has_time:
        raise ValidationFailed("Scoring compares recent months, and this dataset has none.")
    inputs = {f.key: await _factor_input(session, actor, dataset, f, cohort) for f in spec.factors}
    subject = " · ".join(cohort.values()) if cohort else f"All of {dataset.name}"
    try:
        result = score_cohort(spec, inputs, f"Cohort · {subject}")
    except ModelError as error:
        raise ValidationFailed(str(error)) from error
    prediction = Prediction(
        model_version_id=model.id,
        requested_by=actor.member_id,
        subject=result["subject"],
        cohort=cohort,
        score=result["score"],
        band=result["band"],
        factors=result["factors"],
    )
    session.add(prediction)
    await audit.record(
        session,
        actor,
        "Scored prediction",
        "model",
        model.id,
        spec.label,
        f"{result['subject']}: {result['score']} ({result['band']})",
    )
    await session.commit()
    return prediction_dict(prediction, model)


async def history(session: AsyncSession, actor: Actor, limit: int = 20) -> list[dict[str, Any]]:
    rows = await session.execute(
        select(Prediction, ModelVersion)
        .join(ModelVersion, ModelVersion.id == Prediction.model_version_id)
        .where(Prediction.requested_by == actor.id)
        .order_by(Prediction.created_at.desc())
        .limit(limit)
    )
    return [prediction_dict(p, m) for p, m in rows]


async def _owned(session: AsyncSession, actor: Actor, prediction_id: str) -> tuple[Prediction, ModelVersion]:
    prediction = await session.get(Prediction, prediction_id)
    if prediction is None or (
        prediction.requested_by != actor.id and not actor.can(Permission.VIEW_ALL_AUDIT)
    ):
        raise NotFound("Prediction not found.")
    model = await session.get(ModelVersion, prediction.model_version_id)
    assert model is not None
    return prediction, model


def _pdf(data: dict[str, Any]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_fill_color(245, 245, 235)
    pdf.rect(0, 0, 210, 297, "F")
    pdf.set_text_color(55, 61, 64)
    pdf.set_font("Helvetica", "B", 20)
    pdf.cell(0, 12, "DataForge prediction", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.set_text_color(82, 88, 91)
    pdf.multi_cell(
        0,
        6,
        f"{data['subject']}\n{data['outcome']}\nModel {data['modelVersion']} · {data['createdAt'][:16].replace('T', ' ')} UTC",
    )
    pdf.ln(4)
    pdf.set_text_color(233, 80, 54)
    pdf.set_font("Helvetica", "B", 40)
    pdf.cell(40, 18, str(data["score"]))
    pdf.set_text_color(55, 61, 64)
    pdf.set_font("Helvetica", "", 14)
    pdf.cell(0, 18, f"/ 100 · {data['band']}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Factors, largest influence first", new_x="LMARGIN", new_y="NEXT")
    for factor in data["factors"]:
        pdf.set_font("Helvetica", "B", 11)
        arrow = "raises" if factor["direction"] == "up" else "lowers"
        pdf.cell(
            0,
            7,
            f"{factor['name']}  (weight {factor['weight']:.2f}, {arrow} the score)",
            new_x="LMARGIN",
            new_y="NEXT",
        )
        pdf.set_font("Helvetica", "", 10)
        pdf.multi_cell(0, 5, factor["reading"])
        pdf.ln(1)
    pdf.ln(3)
    pdf.set_font("Helvetica", "I", 9)
    pdf.set_text_color(107, 113, 117)
    pdf.multi_cell(
        0,
        5,
        "Scores describe cohorts, not individual people. A score is a likelihood, not a certainty. "
        "The factor weights are documented with the model version and add up to 1.",
    )
    return bytes(pdf.output())


def _xlsx(data: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "Prediction"
    header = Font(bold=True, color="373D40")
    fill = PatternFill("solid", fgColor="EBEBD9")
    for row in [
        ("Subject", data["subject"]),
        ("Outcome", data["outcome"]),
        ("Score", data["score"]),
        ("Band", data["band"]),
        ("Model version", data["modelVersion"]),
        ("Scored at (UTC)", data["createdAt"]),
    ]:
        sheet.append(row)
        sheet.cell(sheet.max_row, 1).font = header
    sheet.append([])
    sheet.append(["Factor", "Weight", "Direction", "Contribution", "Reading"])
    for cell in sheet[sheet.max_row]:
        cell.font, cell.fill = header, fill
    for factor in data["factors"]:
        sheet.append(
            [
                factor["name"],
                factor["weight"],
                factor["direction"],
                factor.get("contribution"),
                factor["reading"],
            ]
        )
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["E"].width = 90
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def export(session: AsyncSession, actor: Actor, prediction_id: str, fmt: str) -> tuple[bytes, str, str]:
    prediction, model = await _owned(session, actor, prediction_id)
    data = prediction_dict(prediction, model)
    if fmt == "pdf":
        content, media, extension = _pdf(data), "application/pdf", "pdf"
    elif fmt == "xlsx":
        content, media, extension = (
            _xlsx(data),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        )
    else:
        raise ValidationFailed("Export as pdf or xlsx.")
    await audit.record(
        session,
        actor,
        "Exported prediction",
        "prediction",
        prediction.id,
        data["subject"],
        f"As {extension.upper()}",
    )
    await session.commit()
    return content, media, f"dataforge-prediction-{prediction.id[:8]}.{extension}"
